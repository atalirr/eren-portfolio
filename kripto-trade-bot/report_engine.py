import json
import os
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import state_manager

LEDGER_FILE = "bot_ledger.json"
REPORTS_DIR = "reports"

def _fmt_try(val, decimals=2):
    """TL formatında sayı döndürür (nokta=binlik, virgül=ondalık)."""
    abs_val = abs(val)
    if decimals == 0:
        formatted = f"{abs_val:,.0f}".replace(",", ".")
    else:
        formatted = f"{abs_val:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return formatted

def generate_daily_report(target_date=None):
    """
    Belirli bir gün için özet rapor oluşturur ve JSON olarak kaydeder.
    target_date: 'YYYY-MM-DD' formatında string. None ise dünün tarihini alır.
    """
    if not os.path.exists(LEDGER_FILE):
        return None

    if target_date is None:
        target_date = (state_manager.get_tr_now() - timedelta(days=1)).strftime("%Y-%m-%d")

    try:
        # ── v3.1: trade_history.jsonl'den oku (asıl kaynak) ──
        # bot_ledger.json'daki trade_log son 100 ile sınırlı,
        # trade_history.jsonl TÜM işlemleri içerir.
        trade_log = []
        history_file = "trade_history.jsonl"
        if os.path.exists(history_file):
            with open(history_file, "r", encoding="utf-8") as hf:
                for line in hf:
                    line = line.strip()
                    if line:
                        try:
                            trade_log.append(json.loads(line))
                        except Exception:
                            pass
        
        # Fallback: trade_history.jsonl yoksa bot_ledger.json'dan oku
        if not trade_log:
            with open(LEDGER_FILE, "r") as f:
                data = json.load(f)
            trade_log = data.get("trade_log", [])
        
        daily_trades = [t for t in trade_log if t.get("date") == target_date]
        
        if not daily_trades:
            return f"📅 {target_date} tarihinde hiç işlem yapılmadı."

        # USDT/TRY kuru
        import trade_executor
        rate = trade_executor.get_usdt_try_rate()

        buys = [t for t in daily_trades if t.get("side") == "BUY"]
        sells = [t for t in daily_trades if t.get("side") in ["SELL", "SELL-ALL", "PARTIAL-SELL"]]
        
        total_pnl_usd = sum([t.get("pnl", 0.0) for t in sells])
        total_pnl_try = total_pnl_usd * rate
        wins = [t for t in sells if t.get("pnl", 0) > 0]
        losses = [t for t in sells if t.get("pnl", 0) <= 0]
        
        win_rate = (len(wins) / len(sells) * 100) if sells else 0
        
        best_trade = max(sells, key=lambda x: x.get("pnl", 0), default=None)
        worst_trade = min(sells, key=lambda x: x.get("pnl", 0), default=None)

        # Toplam portföy değeri
        start_bal = state_manager.db.get("start_balance", 50000.0)
        realized_usd = state_manager.db.get("total_pnl", 0.0)
        portfolio_try = start_bal + (realized_usd * rate)

        # Kumbara
        vault_usd = state_manager.db.get("profit_vault", 0.0)
        vault_try = vault_usd * rate

        report = {
            "date": target_date,
            "total_pnl_usd": round(total_pnl_usd, 2),
            "total_pnl_try": round(total_pnl_try, 2),
            "trades_opened": len(buys),
            "trades_closed": len(sells),
            "win_rate": round(win_rate, 1),
            "win_count": len(wins),
            "loss_count": len(losses),
            "best_trade": best_trade,
            "worst_trade": worst_trade,
            "all_daily_trades": daily_trades,
            "generated_at": state_manager.get_tr_now().strftime("%Y-%m-%d %H:%M:%S")
        }

        # Raporu kaydet
        if not os.path.exists(REPORTS_DIR):
            os.makedirs(REPORTS_DIR)
        
        report_path = os.path.join(REPORTS_DIR, f"report_{target_date}.json")
        with open(report_path, "w") as f:
            json.dump(report, f, indent=4)

        # --- WORD (.docx) RAPOR ---
        docx_path = os.path.join(REPORTS_DIR, f"report_{target_date}.docx")
        doc = Document()
        
        title = doc.add_heading(f'ANTIGRAVITY AI - GÜNLÜK İŞLEM RAPORU', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f'Tarih: {target_date}').alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_heading('1. Genel Performans Özeti', level=1)
        summary_table = doc.add_table(rows=1, cols=2)
        summary_table.style = 'Table Grid'
        hdr_cells = summary_table.rows[0].cells
        hdr_cells[0].text = 'Metrik'
        hdr_cells[1].text = 'Değer'
        
        pnl_sign = "+" if total_pnl_try >= 0 else ""
        summary_items = [
            ('Günlük Net K/Z', f'{pnl_sign}{_fmt_try(total_pnl_try)} ₺'),
            ('Başarı Oranı (Win Rate)', f'%{win_rate:.1f}'),
            ('Açılan Pozisyon', str(len(buys))),
            ('Kapanan Pozisyon', str(len(sells))),
            ('Kârlı İşlem', str(len(wins))),
            ('Zararlı İşlem', str(len(losses))),
            ('Portföy Değeri', f'{_fmt_try(portfolio_try)} ₺'),
            ('Kumbara Birikimi', f'{_fmt_try(vault_try)} ₺'),
        ]
        for metric, value in summary_items:
            row = summary_table.add_row().cells
            row[0].text = metric
            row[1].text = value

        doc.add_heading('2. Detaylı İşlem Geçmişi', level=1)
        trades_table = doc.add_table(rows=1, cols=7)
        trades_table.style = 'Table Grid'
        
        headers = ['Saat', 'Sembol', 'Yön', 'Fiyat', 'Miktar', 'K/Z (₺)', 'Sebep']
        hdr_cells = trades_table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            
        sorted_trades = sorted(daily_trades, key=lambda x: x.get("time", ""))
        for t in sorted_trades:
            pnl_val = t.get('pnl', 0)
            pnl_try_val = pnl_val * rate
            if t.get('side') in ['SELL', 'SELL-ALL', 'PARTIAL-SELL']:
                s = "+" if pnl_try_val >= 0 else ""
                pnl_str = f"{s}{_fmt_try(pnl_try_val)} ₺"
            else:
                pnl_str = "-"
            
            row_cells = trades_table.add_row().cells
            row_cells[0].text = t.get('time', '')
            row_cells[1].text = t.get('symbol', '')
            row_cells[2].text = t.get('side', '')
            row_cells[3].text = str(t.get('price', ''))
            row_cells[4].text = f"{t.get('qty', 0):.4f}"
            row_cells[5].text = pnl_str
            row_cells[6].text = t.get('reason', '-')

        doc.add_paragraph(f'\nBu rapor ANTIGRAVITY AI tarafından otomatik olarak üretilmiştir.')
        doc.add_paragraph(f'Oluşturulma Zamanı: {report["generated_at"]}')
        
        doc.save(docx_path)

        # ── TELEGRAM MESAJI — Modern Çerçevesiz Monospace ──
        pnl_emoji = "🟢" if total_pnl_try >= 0 else "🔴"
        pnl_sign = "+" if total_pnl_try >= 0 else ""

        msg = f"📊 *GÜN SONU RAPORU — {target_date}*\n\n"

        # İşlem Özeti
        msg += "💼 *İŞLEM ÖZETİ*\n"
        msg += "```\n"
        msg += f"• Açılan Pozisyon  : {len(buys)} adet\n"
        msg += f"• Kapanan Pozisyon : {len(sells)} adet\n"
        msg += f"• Başarı Oranı     : %{win_rate:.1f}\n"
        msg += f"• Kârlı İşlem      : ✅ {len(wins)} adet\n"
        msg += f"• Zararlı İşlem    : ❌ {len(losses)} adet\n"
        msg += "```\n\n"

        # Finansal Özet
        msg += "💰 *FİNANSAL ÖZET*\n"
        msg += "```\n"
        msg += f"• Günlük Net K/Z   : {pnl_emoji} {pnl_sign}{_fmt_try(total_pnl_try)} ₺\n"
        msg += f"• Portföy Değeri   : {_fmt_try(portfolio_try)} ₺\n"
        msg += f"• Kumbara Birikimi : 🏦 {_fmt_try(vault_try)} ₺\n"
        msg += "```\n\n"

        # En iyi / En kötü işlem
        if best_trade:
            best_sym = str(best_trade.get('symbol', '')).split('/')[0]
            best_pnl_try = best_trade.get('pnl', 0) * rate
            best_sign = "+" if best_pnl_try >= 0 else ""
            msg += f"🏆 *En İyi İşlem:* `{best_sym}` → `{best_sign}{_fmt_try(best_pnl_try)} ₺`\n"
        if worst_trade:
            worst_sym = str(worst_trade.get('symbol', '')).split('/')[0]
            worst_pnl_try = worst_trade.get('pnl', 0) * rate
            worst_sign = "+" if worst_pnl_try >= 0 else ""
            msg += f"💀 *En Kötü İşlem:* `{worst_sym}` → `{worst_sign}{_fmt_try(worst_pnl_try)} ₺`\n"

        # İşlem detayları
        if sells:
            msg += "\n📜 *İŞLEM DETAYLARI*\n"
            msg += "```\n"
            for s in sorted(sells, key=lambda x: x.get("time", "")):
                sym = str(s.get("symbol", "")).split("/")[0][:5]
                t_str = str(s.get("time", ""))[:5]
                pnl_usd = s.get("pnl", 0) or 0
                p_try = pnl_usd * rate
                emoji = "🟢" if pnl_usd >= 0 else "🔴"
                s_sign = "+" if pnl_usd >= 0 else ""
                msg += f"{emoji} {sym:<5} {t_str}  {s_sign}{_fmt_try(p_try, 0)} ₺\n"
            msg += "```\n"

        msg += f"\n📂 _Word raporu kaydedildi: report\\_{target_date}.docx_"
        
        return msg

    except Exception as e:
        return f"❌ Rapor oluşturulurken hata: {e}"

if __name__ == "__main__":
    print(generate_daily_report('2026-05-05'))
